import streamlit as st
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from webdriver_manager.core.os_manager import ChromeType
from selenium.webdriver.common.by import By
import time
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
import io

def setup_bot():
    options = webdriver.ChromeOptions()
    options.add_argument('--headless=new')
    options.add_argument('--no-sandbox')
    options.add_argument('--disable-dev-shm-usage')
    options.add_argument('--disable-gpu')

    # Webdriver Manager ile Chromium'a uygun sürümü otomatik indir ve kullan
    service = Service(ChromeDriverManager(chrome_type=ChromeType.CHROMIUM).install())
    driver = webdriver.Chrome(service=service, options=options)
    return driver

st.set_page_config(page_title="Araç Kiralama Botu", page_icon="🚗", layout="wide")
st.title("🚗 Araç Kiralama Fiyat Toplama Botu")
st.markdown("Bu uygulama **Yolcu360** ve **Obilet** sitelerindeki güncel fiyatları seçtiğiniz lokasyona göre çeker.")

# --- LOKASYON LİNKLERİ SÖZLÜĞÜ ---
LOKASYONLAR = {
    "Istanbul Havalimanı": {
        "Yolcu360": "https://yolcu360.com/arac-kiralama/search?a=30-65&p_d=2026-10-10&d_d=2026-10-11&p_t=10:00&d_t=10:00&sb=recommended&p_p=41.274874,28.732136&pid=ChIJqZW8Cvb_n0ARBuUkyCzgDDg",
        "Obilet": "https://www.obilet.com/araclar/24-24/20261010_1000-20261011_1000"
    },
    "Adnan Menderes Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Ankara Çankaya": {"Yolcu360": "", "Obilet": ""},
    "Ankara Esenboğa": {"Yolcu360": "", "Obilet": ""},
    "Ankara Esenboğa Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Antalya": {"Yolcu360": "", "Obilet": ""},
    "Ataşehir": {"Yolcu360": "", "Obilet": ""},
    "Denizli Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Denizli Pamukkale": {"Yolcu360": "", "Obilet": ""},
    "Diyarbakir Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Erzurum Şehir": {"Yolcu360": "", "Obilet": ""},
    "Kartal": {"Yolcu360": "", "Obilet": ""},
    "Konya Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Sabiha Gökçen Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Samsun Havalimanı": {"Yolcu360": "", "Obilet": ""},
    "Tokat Şehir": {"Yolcu360": "", "Obilet": ""},
    "Ümraniye": {"Yolcu360": "", "Obilet": ""},
    "Yenibosna": {"Yolcu360": "", "Obilet": ""}
}

# --- KULLANICI SEÇİMLERİ ---
col1, col2 = st.columns(2)
with col1:
    secilen_broker = st.selectbox("🔌 Broker Seçimi:", ["Tümü", "Yolcu360", "Obilet"])
with col2:
    secilen_lokasyon = st.selectbox("📍 Veri Çekilecek İstasyonu Seçin:", list(LOKASYONLAR.keys()))

# Sadece Greenmotion Getir Filtresi
sadece_greenmotion = st.checkbox("💚 Sadece Greenmotion Araçlarını Getir", value=False)

if st.button("▶️ Verileri Çek ve Excel Oluştur"):
    all_urls = LOKASYONLAR[secilen_lokasyon]

    # Hangi sitelerin taranacağını belirle
    siteler_to_scan = []
    if secilen_broker == "Tümü":
        siteler_to_scan = ["Yolcu360", "Obilet"]
    else:
        siteler_to_scan = [secilen_broker]

    eksik_url_var = False
    for site in siteler_to_scan:
        if not all_urls[site]:
            st.error(f"⚠️ {secilen_lokasyon} için {site} arama linki eksik!")
            eksik_url_var = True

    if not eksik_url_var:
        with st.spinner(f"{secilen_lokasyon} için bot çalışıyor (Bu işlem 1-2 dakika sürebilir)..."):
            live_data = []
            bot = None
            try:
                bot = setup_bot()
                bot.maximize_window()

                for site_name in siteler_to_scan:
                    url = all_urls[site_name]
                    st.info(f"🌐 {site_name} taranıyor...")
                    bot.get(url)
                    time.sleep(5)

                    # Sayfa aşağı kaydırma
                    last_height = bot.execute_script("return document.body.scrollHeight")
                    while True:
                        bot.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                        time.sleep(3)
                        new_height = bot.execute_script("return document.body.scrollHeight")
                        if new_height == last_height:
                            break
                        last_height = new_height

                    if site_name == "Yolcu360":
                        car_cards = bot.find_elements(By.XPATH, "//div[contains(@class,'car-card')]")
                        for card in car_cards:
                            try:
                                sirket_el = card.find_elements(By.XPATH, './/button[contains(@class,"review-island")]//img[@alt]')
                                if not sirket_el: sirket_el = card.find_elements(By.XPATH, './/img[@alt]')
                                sirket_text = sirket_el[0].get_attribute("alt").strip() if sirket_el else "Bilinmiyor"

                                # Filtreleme: Eğer Sadece Greenmotion seçildiyse ve şirket adı eşleşmiyorsa atla
                                if sadece_greenmotion and "greenmotion" not in sirket_text.lower():
                                    continue

                                model_el = card.find_elements(By.XPATH, ".//div[contains(@class,'text-dark-gray') and contains(@class,'text-lg') and contains(@class,'font-bold')]")
                                model_text = model_el[0].get_attribute("textContent").strip() if model_el else "Bilinmiyor"

                                sinif_el = card.find_elements(By.XPATH, ".//*[@data-cms-key and starts-with(@data-cms-key,'filter_class_type_')]")
                                sinif_text = sinif_el[0].get_attribute("textContent").strip() if sinif_el else "Bilinmiyor"

                                fiyat_el = card.find_elements(By.XPATH, './/*[@id="car_total_price"]')
                                if not fiyat_el: fiyat_el = card.find_elements(By.XPATH, './/*[@data-cms-key="text_daily_price2"]')
                                fiyat_text = fiyat_el[0].get_attribute("textContent").strip() if fiyat_el else ""

                                if fiyat_text and any(char.isdigit() for char in fiyat_text):
                                    live_data.append([site_name, secilen_lokasyon, sirket_text, model_text, sinif_text, fiyat_text.replace('\n', ' ')])
                            except Exception: continue

                    elif site_name == "Obilet":
                        car_cards = bot.find_elements(By.XPATH, "//div[contains(@class,'car-item')]")
                        for card in car_cards:
                            try:
                                sirket_el = card.find_elements(By.XPATH, './/img[contains(@class,"car-provider__logo")]')
                                sirket_text = sirket_el[0].get_attribute("alt").strip() if sirket_el else "Bilinmiyor"

                                # Filtreleme: Eğer Sadece Greenmotion seçildiyse ve şirket adı eşleşmiyorsa atla
                                if sadece_greenmotion and "greenmotion" not in sirket_text.lower():
                                    continue

                                model_el = card.find_elements(By.XPATH, ".//div[contains(@class,'car-item__name')]")
                                model_text = model_el[0].get_attribute("textContent").strip() if model_el else "Bilinmiyor"

                                sinif_el = card.find_elements(By.XPATH, ".//div[contains(@class,'car-item__segment')]")
                                sinif_text = sinif_el[0].get_attribute("textContent").strip() if sinif_el else "Bilinmiyor"

                                fiyat_el = card.find_elements(By.XPATH, ".//div[contains(@class,'car-price')]//div[contains(@class,'badge')]//span[contains(@class,'fw-500')]")
                                if not fiyat_el: fiyat_el = card.find_elements(By.CSS_SELECTOR, ".car-price__text.is-big")
                                if not fiyat_el: fiyat_el = card.find_elements(By.CSS_SELECTOR, ".car-price > .car-price__text:last-child")
                                fiyat_text = fiyat_el[0].get_attribute("textContent").strip() if fiyat_el else ""

                                if fiyat_text and any(char.isdigit() for char in fiyat_text):
                                    live_data.append([site_name, secilen_lokasyon, sirket_text, model_text, sinif_text, fiyat_text.replace('\n', ' ')])
                            except Exception: continue

                st.success(f"✅ Veri çekme tamamlandı! ({len(live_data)} araç bulundu) Excel dosyası hazırlanıyor...")

                if live_data:
                    df_live = pd.DataFrame(live_data, columns=["Broker", "Lokasyon", "Kiralama Şirketi", "Araç Modeli", "Araç Sınıfı", "Fiyat"])

                    output = io.BytesIO()
                    with pd.ExcelWriter(output, engine='openpyxl') as writer:
                        df_live.to_excel(writer, sheet_name="Canlı Veriler", index=False)

                    wb_live = openpyxl.load_workbook(output)
                    ws_live = wb_live["Canlı Veriler"]

                    header_font = Font(bold=True, color="FFFFFF")
                    header_fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
                    greenmotion_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

                    for cell in ws_live[1]:
                        cell.font = header_font
                        cell.fill = header_fill

                    for col_idx, col in enumerate(ws_live.columns, 1):
                        max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
                        ws_live.column_dimensions[get_column_letter(col_idx)].width = max_length + 2

                    # Greenmotion satırlarını yeşile boyama (3. sütun)
                    for row in ws_live.iter_rows(min_row=2, max_row=ws_live.max_row, min_col=1, max_col=ws_live.max_column):
                        if row[2].value and "greenmotion" in str(row[2].value).lower():
                            for cell in row:
                                cell.fill = greenmotion_fill

                    final_output = io.BytesIO()
                    wb_live.save(final_output)
                    final_output.seek(0)

                    st.download_button(
                        label=f"📥 {secilen_lokasyon} Raporunu İndir",
                        data=final_output,
                        file_name=f"{secilen_lokasyon}_Fiyat_Raporu.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

                    st.dataframe(df_live)
                else:
                    st.warning("⚠️ Veri bulunamadı veya aranan kriterde araç yok.")

            except Exception as e:
                st.error(f"❌ Hata oluştu: {e}")
            finally:
                if bot:
                    bot.quit()
